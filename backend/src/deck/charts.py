"""Native chart/table insertion (RM-8, D3) -- real PowerPoint objects, never images.

`add_chart`/`add_table` produce editable OOXML chart/table parts the same way
PowerPoint's own Insert Chart/Table does -- unlike an AI-generated or
stock-photo image, the numbers stay editable and stay on-theme (the chart
picks up the template's own colour scheme; nothing here sets one).

**Not wired into the LD-8 renderer** (`PLAN-LLM-DECK-PLANNING.md` Phase C):
`deck/catalog.py`'s anchor `purpose` vocabulary has no chart/table variant
yet, so nothing currently calls `insert_chart`/`insert_table`. Kept, per the
plan's module disposition table ("keep"), for when a follow-up task defines
how a design's chart-shaped anchor is catalogued and planned -- see the
Phase C report's findings. `ChartSpec`/`TableSpec` live here now, not in the
deleted `deck/spec.py`, since this module is their only remaining consumer.
"""

from __future__ import annotations

from typing import Literal

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Length
from pydantic import BaseModel, ConfigDict, Field, field_validator


class ChartSeries(BaseModel):
    model_config = ConfigDict(extra="forbid")

    name: str = Field(..., min_length=1)
    values: list[float] = Field(..., min_length=1)


class ChartSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")

    kind: Literal["bar", "column", "line", "pie"] = "bar"
    title: str | None = None
    categories: list[str] = Field(..., min_length=1)
    series: list[ChartSeries] = Field(..., min_length=1)

    @field_validator("series")
    @classmethod
    def _series_match_categories(cls, series: list[ChartSeries], info) -> list[ChartSeries]:
        categories = info.data.get("categories")
        if categories is not None:
            for s in series:
                if len(s.values) != len(categories):
                    raise ValueError(
                        f"series {s.name!r} has {len(s.values)} values but there are "
                        f"{len(categories)} categories -- they must match 1:1"
                    )
        return series


class TableSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")

    headers: list[str] = Field(..., min_length=1)
    rows: list[list[str]] = Field(default_factory=list)

    @field_validator("rows")
    @classmethod
    def _rows_match_header_width(cls, rows: list[list[str]], info) -> list[list[str]]:
        headers = info.data.get("headers")
        if headers is not None:
            for row in rows:
                if len(row) != len(headers):
                    raise ValueError(
                        f"row {row!r} has {len(row)} cells but there are {len(headers)} headers"
                    )
        return rows


_CHART_KIND_MAP = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}


def insert_chart(*, slide, chart_spec: ChartSpec, left: Length, top: Length, width: Length, height: Length) -> None:
    data = CategoryChartData()
    data.categories = chart_spec.categories
    for series in chart_spec.series:
        data.add_series(series.name, series.values)

    chart_type = _CHART_KIND_MAP.get(chart_spec.kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
    graphic_frame = slide.shapes.add_chart(chart_type, left, top, width, height, data)

    chart = graphic_frame.chart
    chart.has_title = bool(chart_spec.title)
    if chart_spec.title:
        chart.chart_title.text_frame.text = chart_spec.title


def insert_table(*, slide, table_spec: TableSpec, left: Length, top: Length, width: Length, height: Length) -> None:
    rows = 1 + len(table_spec.rows)
    cols = len(table_spec.headers)
    graphic_frame = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = graphic_frame.table

    for col, header in enumerate(table_spec.headers):
        table.cell(0, col).text = header
    for row_index, row in enumerate(table_spec.rows, start=1):
        for col, value in enumerate(row):
            table.cell(row_index, col).text = value
