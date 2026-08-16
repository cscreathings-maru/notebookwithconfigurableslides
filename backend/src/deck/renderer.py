"""The deck renderer (LD-8) -- `DeckPlan` + template bytes -> deck bytes.

Pure function: no network, no LLM, no object store. Clones the template's
OWN designed slides (`deck/clone.py`) in plan order and writes text into the
anchors the plan specifies -- **never builds a slide from a layout** (L1).
This is what keeps a corporate template's full design (background art,
logos, decoration) intact where the old layout-based renderer discarded it
(`ASSESSMENT-LLM-DECK-PLANNING.md` §1): nothing is reconstructed here, only
cloned and re-labelled.

Every collaborator is either an argument or a local `python-pptx` call --
same posture as the module this replaces, and for the same reason: a pure
function is what makes a golden-file/determinism test meaningful (G8: same
plan + template -> byte-comparable output).

**Fidelity**: writing an anchor's text via `text_frame.text = ...` replaces
its runs with a single new run inheriting the frame's own (theme/placeholder)
formatting -- it never sets `.font.name`, `.font.size` or `.font.color`
explicitly. That is the same limitation the module this replaces had
(`_fill_bullets`) and is accepted for the same reason (D2): fidelity comes
from filling the template's own shapes, never from choosing how anything
looks.

**Known scope gap**: chart/table anchors are not implemented. The old
`DeckSpec` carried `ChartSpec`/`TableSpec` rendered via `deck/charts.py`;
LD-8's task list names only text-anchor cloning, and `deck/catalog.py`'s
anchor `purpose` vocabulary has no chart/table variant. `deck/charts.py`
stays in the tree, unused, until a follow-up task defines how a design's
chart-shaped anchor is catalogued and planned. Recorded in the Phase C
report as a real, not hypothetical, regression versus the pipeline this
replaces -- not silently dropped.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..core.errors import ValidationError
from .plan import DeckPlan
from .clone import clone_slide, drop_slides

_DESIGN_ID_PREFIX = "slide-"


def _slide_index_from_design_id(design_id: str) -> int | None:
    """Inverse of `deck/catalog.py::design_id_for` -- the one place this
    encoding is parsed back, so a format change only needs updating here and
    there."""
    if not design_id.startswith(_DESIGN_ID_PREFIX):
        return None
    try:
        return int(design_id[len(_DESIGN_ID_PREFIX) :])
    except ValueError:
        return None


def _find_shape_by_id(shapes, shape_id: int):
    """Resolves an anchor's shape regardless of group nesting depth
    (assessment §6's named risk: "text inside grouped shapes"). Mirrors
    `deck/dump.py`'s recursive walk so the two never disagree about where a
    given `shape_id` lives."""
    for shape in shapes:
        if shape.shape_id == shape_id:
            return shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            found = _find_shape_by_id(shape.shapes, shape_id)
            if found is not None:
                return found
    return None


def _apply_anchor_texts(slide, anchor_texts: dict[str, str]) -> None:
    for raw_anchor_id, text in anchor_texts.items():
        try:
            shape_id = int(raw_anchor_id)
        except ValueError:
            continue
        shape = _find_shape_by_id(slide.shapes, shape_id)
        if shape is None or not getattr(shape, "has_text_frame", False):
            continue
        shape.text_frame.text = text


def render_deck(*, plan: DeckPlan, template_pptx: bytes) -> bytes:
    """Render `plan` into `template_pptx`'s OWN slides, in plan order.

    Each `PlanSlide.design_id` names an original slide index (LD-9 wires
    this from `deck/plan.py`, which only ever names ids drawn from the
    catalog it was given -- see that module's validation). A design may be
    referenced any number of times (L7, L6's `(lanjutan)` splitting); each
    use clones the SAME source slide independently, so N uses of one design
    yield N intact copies (`deck/clone.py`'s own guarantee).

    Raises `ValidationError` -- never renders a partial/wrong deck -- if a
    plan names a design this template's slide count cannot resolve. That
    should be unreachable given `deck/plan.py`'s own validation against the
    same catalog the plan was built from; kept as a stated invariant here
    too, in case a caller ever assembles a `DeckPlan` by hand.
    """
    prs = Presentation(io.BytesIO(template_pptx))
    original_slides = list(prs.slides)

    new_indexes: list[int] = []
    for slide_spec in plan.slides:
        source_index = _slide_index_from_design_id(slide_spec.design_id)
        if source_index is None or not (0 <= source_index < len(original_slides)):
            raise ValidationError(
                f"Plan references design {slide_spec.design_id!r}, which does not "
                f"resolve to a slide in this template ({len(original_slides)} slides)."
            )
        source_slide = original_slides[source_index]
        new_slide = clone_slide(prs, source_slide)
        _apply_anchor_texts(new_slide, slide_spec.anchor_texts)
        new_indexes.append(len(prs.slides) - 1)

    # Clones are appended after every original slide, so `new_indexes` is
    # already in ascending, plan-matching order -- dropping the originals
    # leaves exactly the rendered deck, in the order the plan specified.
    drop_slides(prs, set(new_indexes))

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
