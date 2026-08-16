"""Unit: the clone-based deck renderer (LD-8) against the REAL BRI template.

G1/G3/G5/G8 from the plan's exit gate, exercised locally: the rendered deck
keeps the template's design (logo, static shapes) intact, a cloned slide with
images stays readable, one design used N times renders N intact copies, and
rendering is deterministic (same plan + template -> byte-comparable output).
"""

from __future__ import annotations

import io

import pytest
from pptx import Presentation

from src.core.errors import ValidationError
from src.deck.catalog import catalog_template, design_id_for
from src.deck.dump import dump_presentation
from src.deck.plan import DeckPlan, PlanSlide
from src.deck.renderer import render_deck
from src.registry.house_template import HOUSE_TEMPLATE_PATH
from tests.fakes import FakeLlm


def _is_picture(shape) -> bool:
    return shape.shape_type is not None and "PICTURE" in str(shape.shape_type)


@pytest.fixture(scope="module")
def template_bytes() -> bytes:
    if not HOUSE_TEMPLATE_PATH.exists():
        pytest.fail(f"House template missing: {HOUSE_TEMPLATE_PATH}")
    return HOUSE_TEMPLATE_PATH.read_bytes()


@pytest.fixture(scope="module")
async def bri_catalog(template_bytes: bytes):
    dumps = dump_presentation(template_bytes)
    catalog, _usage = await catalog_template(dumps=dumps, llm=FakeLlm(), provider_config={})
    return catalog


def _open(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


def test_rendering_the_cover_design_keeps_its_logo(template_bytes: bytes, bri_catalog) -> None:
    """G1 -- the failure that triggered this whole plan: a rendered deck must
    keep the BRI logo (slide 0's picture), not discard it the way the old
    layout-based renderer did."""
    cover = bri_catalog.by_id(design_id_for(0))
    assert cover is not None and cover.usable
    title_anchor = next(a for a in cover.anchors if a.purpose == "title")

    plan = DeckPlan(slides=[PlanSlide(design_id=cover.design_id, anchor_texts={title_anchor.anchor_id: "Judul Baru"})])
    deck_bytes = render_deck(plan=plan, template_pptx=template_bytes)

    rendered = _open(deck_bytes)
    assert len(rendered.slides) == 1
    pictures = [s for s in rendered.slides[0].shapes if _is_picture(s)]
    assert pictures, "the cover's logo must survive rendering"
    for picture in pictures:
        assert picture.image.blob, "cloned image must resolve to real bytes (SC-0)"


def test_anchor_text_is_written_into_the_right_shape(template_bytes: bytes, bri_catalog) -> None:
    cover = bri_catalog.by_id(design_id_for(0))
    title_anchor = next(a for a in cover.anchors if a.purpose == "title")

    plan = DeckPlan(slides=[PlanSlide(design_id=cover.design_id, anchor_texts={title_anchor.anchor_id: "Laporan Q4 2026"})])
    rendered = _open(render_deck(plan=plan, template_pptx=template_bytes))

    shape = next(s for s in rendered.slides[0].shapes if s.shape_id == int(title_anchor.anchor_id))
    assert shape.text_frame.text == "Laporan Q4 2026"


def test_richest_slide_design_keeps_every_shape(template_bytes: bytes, bri_catalog) -> None:
    """Slide 20 (the timeline): 95 shapes, 3 images (assessment §1) -- must
    render with none silently dropped, same guarantee `deck/clone.py`'s own
    suite proves at the primitive level."""
    timeline = bri_catalog.by_id(design_id_for(20))
    assert timeline is not None

    plan = DeckPlan(
        slides=[
            PlanSlide(
                design_id=timeline.design_id,
                anchor_texts={a.anchor_id: "x" for a in timeline.anchors},
            )
        ]
    )
    rendered = _open(render_deck(plan=plan, template_pptx=template_bytes))
    assert len(rendered.slides[0].shapes) == 95
    pictures = [s for s in rendered.slides[0].shapes if _is_picture(s)]
    assert len(pictures) == 3
    for picture in pictures:
        assert picture.image.blob


def test_one_design_used_four_times_renders_four_intact_copies(template_bytes: bytes, bri_catalog) -> None:
    """G5 -- a design may be reused any number of times (L7)."""
    cover = bri_catalog.by_id(design_id_for(0))
    title_anchor = next(a for a in cover.anchors if a.purpose == "title")

    plan = DeckPlan(
        slides=[
            PlanSlide(design_id=cover.design_id, anchor_texts={title_anchor.anchor_id: f"Slide {i}"})
            for i in range(4)
        ]
    )
    rendered = _open(render_deck(plan=plan, template_pptx=template_bytes))
    assert len(rendered.slides) == 4
    for i, slide in enumerate(rendered.slides):
        shape = next(s for s in slide.shapes if s.shape_id == int(title_anchor.anchor_id))
        assert shape.text_frame.text == f"Slide {i}"
        assert any(_is_picture(s) for s in slide.shapes)


def test_render_order_matches_plan_order_not_original_slide_order(template_bytes: bytes, bri_catalog) -> None:
    cover = bri_catalog.by_id(design_id_for(0))
    timeline = bri_catalog.by_id(design_id_for(20))
    title_anchor = next(a for a in cover.anchors if a.purpose == "title")

    # Plan order: timeline THEN cover -- the reverse of their original
    # positions in the template (20, then 0).
    plan = DeckPlan(
        slides=[
            PlanSlide(design_id=timeline.design_id, anchor_texts={a.anchor_id: "x" for a in timeline.anchors[:1]}),
            PlanSlide(design_id=cover.design_id, anchor_texts={title_anchor.anchor_id: "Last"}),
        ]
    )
    rendered = _open(render_deck(plan=plan, template_pptx=template_bytes))
    assert len(rendered.slides) == 2
    assert len(rendered.slides[0].shapes) == 95  # timeline first
    assert any(_is_picture(s) for s in rendered.slides[1].shapes)  # cover second


def test_only_the_planned_slides_survive_not_the_original_30(template_bytes: bytes, bri_catalog) -> None:
    cover = bri_catalog.by_id(design_id_for(0))
    title_anchor = next(a for a in cover.anchors if a.purpose == "title")
    plan = DeckPlan(slides=[PlanSlide(design_id=cover.design_id, anchor_texts={title_anchor.anchor_id: "Only"})])

    rendered = _open(render_deck(plan=plan, template_pptx=template_bytes))
    assert len(rendered.slides) == 1  # not 31 (30 originals + 1 clone)


def test_rendering_is_deterministic(template_bytes: bytes, bri_catalog) -> None:
    """G8: same plan + template -> byte-comparable output."""
    cover = bri_catalog.by_id(design_id_for(0))
    title_anchor = next(a for a in cover.anchors if a.purpose == "title")
    plan = DeckPlan(slides=[PlanSlide(design_id=cover.design_id, anchor_texts={title_anchor.anchor_id: "Stable"})])

    first = render_deck(plan=plan, template_pptx=template_bytes)
    second = render_deck(plan=plan, template_pptx=template_bytes)
    assert first == second


def test_unknown_design_id_raises_rather_than_rendering_a_broken_deck(template_bytes: bytes) -> None:
    plan = DeckPlan(slides=[PlanSlide(design_id="slide-9999", anchor_texts={"1": "x"})])
    with pytest.raises(ValidationError):
        render_deck(plan=plan, template_pptx=template_bytes)


def test_never_sets_explicit_font_overrides(template_bytes: bytes, bri_catalog) -> None:
    """D2 fidelity: text is written by replacing a run, never by setting
    `.font.name`/`.font.size`/`.font.color` -- fidelity comes from the
    template's own theme, not from this module choosing anything."""
    cover = bri_catalog.by_id(design_id_for(0))
    title_anchor = next(a for a in cover.anchors if a.purpose == "title")
    plan = DeckPlan(slides=[PlanSlide(design_id=cover.design_id, anchor_texts={title_anchor.anchor_id: "Themed"})])

    rendered = _open(render_deck(plan=plan, template_pptx=template_bytes))
    shape = next(s for s in rendered.slides[0].shapes if s.shape_id == int(title_anchor.anchor_id))
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            assert run.font.name is None
            assert run.font.size is None
