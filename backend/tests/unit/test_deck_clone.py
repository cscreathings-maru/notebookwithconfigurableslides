"""Unit: image-safe slide cloning (SC-0).

Runs against the REAL BRI template, not a synthetic fixture. That is
deliberate and is the process fix from `PLAN-SLIDE-CLONING.md` §7: the previous
renderer's whole test suite passed against `python-pptx`'s bundled template --
which, unlike a corporate deck, puts its design in layouts -- and shipped a
deck with no logo. A fixture that shares the code's assumption cannot falsify
it.
"""

from __future__ import annotations

import copy
import io

import pytest
from pptx import Presentation

from src.deck.clone import clone_slide, drop_slides, move_slide
from src.registry.house_template import HOUSE_TEMPLATE_PATH


def _is_picture(shape) -> bool:
    return shape.shape_type is not None and "PICTURE" in str(shape.shape_type)


@pytest.fixture(scope="module")
def template_bytes() -> bytes:
    if not HOUSE_TEMPLATE_PATH.exists():
        pytest.fail(f"House template missing: {HOUSE_TEMPLATE_PATH}")
    return HOUSE_TEMPLATE_PATH.read_bytes()


def _open(template_bytes: bytes) -> Presentation:
    return Presentation(io.BytesIO(template_bytes))


def _reopen(prs: Presentation) -> Presentation:
    """Round-trip through bytes. Corruption of this kind only shows on reload --
    an in-memory object still looks fine."""
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


def _slide_with_images(prs: Presentation):
    for slide in prs.slides:
        if sum(1 for s in slide.shapes if _is_picture(s)) >= 2:
            return slide
    pytest.skip("template has no multi-image slide to exercise the remap")


def test_cloned_images_are_readable(template_bytes: bytes) -> None:
    """THE test this module exists for. A naive deepcopy passes every shape
    count below and fails only here, with KeyError: rId3."""
    prs = _open(template_bytes)
    source = _slide_with_images(prs)
    expected = sum(1 for s in source.shapes if _is_picture(s))

    clone_slide(prs, source)
    clone = _reopen(prs).slides[-1]

    pictures = [s for s in clone.shapes if _is_picture(s)]
    assert len(pictures) == expected
    for picture in pictures:
        assert picture.image.blob, "cloned image resolved to no bytes"


def test_naive_deepcopy_would_break_images(template_bytes: bytes) -> None:
    """Pins WHY the remap exists. If a future refactor drops it, the test above
    fails -- and this one proves the failure was not hypothetical."""
    prs = _open(template_bytes)
    source = _slide_with_images(prs)

    naive = prs.slides.add_slide(source.slide_layout)
    for shape in list(naive.shapes):
        shape._element.getparent().remove(shape._element)
    for shape in source.shapes:
        naive.shapes._spTree.append(copy.deepcopy(shape._element))  # no remap

    clone = _reopen(prs).slides[-1]
    with pytest.raises(KeyError):
        for picture in (s for s in clone.shapes if _is_picture(s)):
            _ = picture.image.blob


def test_clone_preserves_every_shape(template_bytes: bytes) -> None:
    prs = _open(template_bytes)
    source = max(prs.slides, key=lambda s: len(s.shapes))  # the richest design
    before = len(source.shapes)
    assert before > 20, "expected a genuinely complex slide to exercise this"

    clone_slide(prs, source)
    assert len(_reopen(prs).slides[-1].shapes) == before


def test_a_design_can_be_cloned_many_times(template_bytes: bytes) -> None:
    """S4: the matcher may reuse one design freely -- a 4-section deck needs
    its divider four times."""
    prs = _open(template_bytes)
    source = _slide_with_images(prs)
    original_count = len(prs.slides)

    for _ in range(4):
        clone_slide(prs, source)

    reopened = _reopen(prs)
    assert len(reopened.slides) == original_count + 4
    for clone in list(reopened.slides)[-4:]:
        for picture in (s for s in clone.shapes if _is_picture(s)):
            assert picture.image.blob


def test_drop_slides_keeps_exactly_the_requested_originals(template_bytes: bytes) -> None:
    prs = _open(template_bytes)
    keep = {0, 2, 20}
    expected_shape_counts = [len(prs.slides[i].shapes) for i in sorted(keep)]

    drop_slides(prs, keep)
    reopened = _reopen(prs)

    assert len(reopened.slides) == len(keep)
    # Identity by shape count: the kept slides must be the ones asked for, in
    # order -- not merely "three slides survived".
    assert [len(s.shapes) for s in reopened.slides] == expected_shape_counts


def test_dropping_preserves_the_design_on_kept_slides(template_bytes: bytes) -> None:
    """The regression that started this plan: rendering must not strip design."""
    prs = _open(template_bytes)
    richest = max(range(len(prs.slides)), key=lambda i: len(prs.slides[i].shapes))
    static_before = sum(1 for s in prs.slides[richest].shapes if not s.is_placeholder)

    drop_slides(prs, {richest})
    kept = _reopen(prs).slides[0]

    assert sum(1 for s in kept.shapes if not s.is_placeholder) == static_before
    assert static_before > 10, "fixture should carry real design furniture"


def test_move_slide_reorders(template_bytes: bytes) -> None:
    prs = _open(template_bytes)
    drop_slides(prs, {0, 2, 20})
    counts = [len(s.shapes) for s in prs.slides]

    move_slide(prs, 2, 0)  # last -> first

    reordered = [len(s.shapes) for s in _reopen(prs).slides]
    assert reordered == [counts[2], counts[0], counts[1]]
