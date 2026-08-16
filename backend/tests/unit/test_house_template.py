"""The bundled house template is real, usable, and renders (RM-15).

A committed binary with no test is a liability: it can be replaced with
something unrenderable and nothing notices until a user tries to generate.
These run against the actual asset, so swapping in a bad `.pptx` fails here
rather than in production.
"""

from __future__ import annotations

import io

import pytest
from pptx import Presentation

from src.deck.catalog import catalog_template
from src.deck.dump import dump_presentation
from src.deck.plan import DeckPlan, PlanSlide
from src.deck.renderer import render_deck
from src.generation.artifact import inspect_pptx
from src.registry.house_template import HOUSE_TEMPLATE_PATH
from tests.fakes import FakeLlm


@pytest.fixture(scope="module")
def house_bytes() -> bytes:
    if not HOUSE_TEMPLATE_PATH.exists():
        pytest.fail(f"House template asset is missing: {HOUSE_TEMPLATE_PATH}")
    return HOUSE_TEMPLATE_PATH.read_bytes()


@pytest.fixture(scope="module")
async def house_catalog(house_bytes: bytes):
    dumps = dump_presentation(house_bytes)
    catalog, _usage = await catalog_template(dumps=dumps, llm=FakeLlm(), provider_config={})
    return catalog


def test_the_asset_dumps_and_catalogues_as_usable(house_catalog) -> None:
    assert house_catalog.usable_designs, "the house template must offer at least one usable design"


def test_it_keeps_its_designed_slides(house_bytes: bytes) -> None:
    """SC-7, reversing RM-15. These slides were stripped to save 14MB -- and
    the 14MB WAS the design: 476 shapes and 57 images live on them, while the
    masters carry no images at all. Rendering clones these slides
    (`deck/clone.py`), so stripping them is what produced a deck with no logo.

    The size is the cost of shipping a real corporate template, not waste."""
    prs = Presentation(io.BytesIO(house_bytes))
    assert len(prs.slides) >= 20, "the designed slides ARE the render vocabulary"

    shapes = sum(len(s.shapes) for s in prs.slides)
    images = sum(
        1
        for s in prs.slides
        for sh in s.shapes
        if sh.shape_type is not None and "PICTURE" in str(sh.shape_type)
    )
    assert shapes > 400, f"expected the full design furniture, found {shapes} shapes"
    assert images > 40, f"expected the brand imagery, found {images} images"


def test_it_is_widescreen(house_bytes: bytes) -> None:
    prs = Presentation(io.BytesIO(house_bytes))
    ratio = prs.slide_width / prs.slide_height
    assert abs(ratio - 16 / 9) < 0.01, f"expected 16:9, got {ratio:.2f}"


def test_a_full_deck_renders_from_it(house_bytes: bytes, house_catalog) -> None:
    """End-to-end against the real asset: a hand-built plan covering the
    cover design plus several content designs, cloned and rendered."""
    usable = house_catalog.usable_designs
    assert len(usable) >= 2, "expected more than one usable design to build a real plan"

    slides: list[PlanSlide] = []
    for design in usable[:4]:
        title_anchor = next((a for a in design.anchors if a.purpose == "title"), None)
        if title_anchor is None:
            continue
        slides.append(PlanSlide(design_id=design.design_id, anchor_texts={title_anchor.anchor_id: "Laporan Kinerja"}))
    assert slides, "expected at least one design with a title anchor"

    plan = DeckPlan(slides=slides)
    deck = render_deck(plan=plan, template_pptx=house_bytes)
    facts = inspect_pptx(deck)
    assert facts.slide_count == len(slides), "the template's own slides must not survive into the deck"
    assert "Laporan Kinerja" in facts.titles or "Laporan Kinerja" in facts.text


def test_rendering_never_overrides_the_brand_font_or_colour(house_bytes: bytes, house_catalog) -> None:
    """D2's invariant, checked against the REAL brand deck rather than a
    synthetic fixture -- an explicit font here would silently discard BRI's
    own typography.

    Scoped to the ANCHOR shape the renderer actually wrote into, not every
    shape on the slide: cloning a real designed slide carries over shapes
    the renderer never touched (e.g. a subtitle placeholder), and THOSE
    legitimately keep whatever font the template's own designer set -- that
    is the template's design, not something this renderer chose. D2 is a
    claim about the renderer's OWN write, not about the source deck's
    pre-existing styling.
    """
    cover = next(d for d in house_catalog.usable_designs if any(a.purpose == "title" for a in d.anchors))
    title_anchor = next(a for a in cover.anchors if a.purpose == "title")
    plan = DeckPlan(slides=[PlanSlide(design_id=cover.design_id, anchor_texts={title_anchor.anchor_id: "Cover"})])
    deck = render_deck(plan=plan, template_pptx=house_bytes)

    rendered_slide = Presentation(io.BytesIO(deck)).slides[0]
    shape = next(s for s in rendered_slide.shapes if s.shape_id == int(title_anchor.anchor_id))

    checked = False
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            checked = True
            assert run.font.name is None
            assert run.font.size is None
            assert run.font.color.type is None
    assert checked, "no runs were produced, so nothing was actually verified"
